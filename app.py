
from pathlib import Path
from datetime import date

import streamlit as st
import pandas as pd
from openpyxl import load_workbook

# =====================================================
# KONFIGURASI APLIKASI
# =====================================================

st.set_page_config(
    page_title="RAB Generator",
    page_icon="📊",
    layout="wide",
)

MASTER_FILE = Path("Master/Master_RAB_Generator.xlsx")

st.title("RAB Generator")
st.caption(
    "V1 — Form data event, upload FK/JK, dan pemeriksaan input"
)


# =====================================================
# MEMBACA MASTER EXCEL
# =====================================================

@st.cache_data
def load_master_data(file_path, modified_time):
    if not Path(file_path).exists():
        raise FileNotFoundError(
            f"File master tidak ditemukan: {file_path}"
        )

    wb = load_workbook(file_path, data_only=True)

    required_sheets = [
        "TEMPLATE RAB",
        "INDEKS RAB",
        "COST CENTER & BEBAN ANGGARAN",
        "NORMA BIAYA",
    ]

    missing = [s for s in required_sheets if s not in wb.sheetnames]

    if missing:
        raise ValueError(
            "Sheet berikut tidak ditemukan: " + ", ".join(missing)
        )

    ws_norma = wb["NORMA BIAYA"]
    ws_indeks = wb["INDEKS RAB"]
    ws_cc = wb["COST CENTER & BEBAN ANGGARAN"]

    # Kategori yang sudah teridentifikasi pada kolom B
    categories = sorted({
        str(row[0]).strip()
        for row in ws_norma.iter_rows(
            min_row=2, min_col=2, max_col=2, values_only=True
        )
        if row[0] is not None
        and str(row[0]).strip() in {"R1", "R2", "L1", "L2"}
    })

    # Indeks lokasi: kolom G = nama/lokasi, H = indeks
    locations = []
    for row in ws_indeks.iter_rows(
        min_row=2, min_col=7, max_col=8, values_only=True
    ):
        name, index = row
        if name is not None and str(name).strip():
            locations.append({
                "name": str(name).strip(),
                "index": index,
            })

    # Cost Center: kolom B = kode, C = nama/keterangan
    cost_centers = []
    for row in ws_cc.iter_rows(
        min_row=2, min_col=2, max_col=3, values_only=True
    ):
        code, description = row
        if code is not None and str(code).strip():
            cost_centers.append({
                "code": str(code).strip(),
                "description": (
                    str(description).strip()
                    if description is not None else ""
                ),
            })

    # Akun beban: kolom E = kode, F = nama/keterangan
    accounts = []
    for row in ws_cc.iter_rows(
        min_row=2, min_col=5, max_col=6, values_only=True
    ):
        code, description = row
        if code is not None and str(code).strip():
            accounts.append({
                "code": str(code).strip(),
                "description": (
                    str(description).strip()
                    if description is not None else ""
                ),
            })

    return categories, locations, cost_centers, accounts


try:
    if not MASTER_FILE.exists():
        st.error(
            "File master belum tersedia. Pastikan file "
            "`Master_RAB_Generator.xlsx` berada di folder `Master/`."
        )
        st.stop()

    categories, locations, cost_centers, accounts = load_master_data(
        str(MASTER_FILE),
        MASTER_FILE.stat().st_mtime,
    )

except Exception as exc:
    st.error(f"Gagal membaca file master: {exc}")
    st.stop()


if not categories:
    st.warning(
        "Kategori R1/R2/L1/L2 belum terbaca. "
        "Periksa posisi data pada sheet NORMA BIAYA."
    )

if not locations:
    st.warning(
        "Daftar lokasi belum terbaca. "
        "Periksa kolom G-H pada sheet INDEKS RAB."
    )


# =====================================================
# FORM DATA EVENT
# =====================================================

st.header("1. Data Event")

# Semua input dan tombol submit harus berada di dalam form yang sama.
with st.form("rab_form"):
    col1, col2 = st.columns(2)

    with col1:
        event_id = st.text_input(
            "Event ID *",
            placeholder="Masukkan ID event",
        )

        kategori = st.selectbox(
            "Kategori Event *",
            options=categories,
            index=None,
            placeholder="Pilih kategori",
        )

        akomodasi = st.selectbox(
            "Akomodasi *",
            options=[
                "Ditanggung Penyelenggara",
                "Tidak Ditanggung Penyelenggara",
            ],
            index=None,
            placeholder="Pilih status akomodasi",
        )

        location_labels = [
            (
                f"{item['name']} — Indeks: {item['index']}"
                if item["index"] is not None
                else item["name"]
            )
            for item in locations
        ]

        selected_location = st.selectbox(
            "Indeks Lokasi *",
            options=location_labels,
            index=None,
            placeholder="Pilih lokasi",
        )

        tanggal_mulai = st.date_input(
            "Tanggal Pelaksanaan *",
            value=None,
        )

    with col2:
        tanggal_selesai = st.date_input(
            "Tanggal Selesai *",
            value=None,
        )

        lokasi_event = st.text_input(
            "Lokasi Event *",
            placeholder="Kota atau lokasi kegiatan",
        )

        cc_labels = [
            f"{item['code']} — {item['description']}"
            for item in cost_centers
        ]

        selected_cc = st.selectbox(
            "Cost Center *",
            options=cc_labels,
            index=None,
            placeholder="Pilih Cost Center",
        )

        account_labels = [
            f"{item['code']} — {item['description']}"
            for item in accounts
        ]

        selected_account = st.selectbox(
            "Akun Beban Anggaran *",
            options=account_labels,
            index=None,
            placeholder="Pilih akun anggaran",
        )

    st.divider()
    st.subheader("2. Upload Dokumen FK/JK")

    uploaded_file = st.file_uploader(
        "Pilih dokumen FK/JK",
        type=["pdf", "png", "jpg", "jpeg"],
        help="Format yang didukung: PDF, PNG, JPG, dan JPEG.",
    )

    submitted = st.form_submit_button(
        "Create RAB",
        type="primary",
        use_container_width=True,
    )


# =====================================================
# TABEL REVIEW FK/JK
# =====================================================

st.header("3. Review Rincian FK/JK")
st.caption(
    "Periksa dan lengkapi rincian kegiatan dari dokumen FK/JK. "
    "Isi tabel ini secara manual; versi ini belum mengekstrak tabel dari PDF/gambar secara otomatis."
)

if uploaded_file is not None:
    st.info(
        f"Dokumen terunggah: **{uploaded_file.name}**. "
        "Gunakan dokumen tersebut sebagai acuan saat mengisi tabel review di bawah."
    )

review_columns = [
    "Jenis Kegiatan",
    "Volume",
    "Satuan",
    "Harga Satuan (Rp)",
    "Keterangan",
]

if "review_fk_jk_data" not in st.session_state:
    st.session_state["review_fk_jk_data"] = pd.DataFrame([
        {
            "Jenis Kegiatan": "",
            "Volume": 0.0,
            "Satuan": "",
            "Harga Satuan (Rp)": 0.0,
            "Keterangan": "",
        }
    ])

edited_review = st.data_editor(
    st.session_state["review_fk_jk_data"],
    key="review_fk_jk_editor",
    num_rows="dynamic",
    use_container_width=True,
    hide_index=True,
    column_config={
        "Jenis Kegiatan": st.column_config.TextColumn(
            "Jenis Kegiatan", help="Nama kegiatan/barang/jasa sesuai FK/JK", required=True
        ),
        "Volume": st.column_config.NumberColumn(
            "Volume", min_value=0.0, step=1.0, format="%.2f", required=True
        ),
        "Satuan": st.column_config.TextColumn(
            "Satuan", help="Contoh: orang, hari, paket"
        ),
        "Harga Satuan (Rp)": st.column_config.NumberColumn(
            "Harga Satuan (Rp)", min_value=0.0, step=1000.0, format="%.0f"
        ),
        "Keterangan": st.column_config.TextColumn("Keterangan"),
    },
)

# Simpan perubahan tabel agar tetap tersedia saat aplikasi melakukan rerun.
st.session_state["review_fk_jk_data"] = edited_review.copy()

valid_review = edited_review[
    edited_review["Jenis Kegiatan"].fillna("").astype(str).str.strip().ne("")
].copy()
if not valid_review.empty:
    valid_review["Estimasi Total (Rp)"] = (
        pd.to_numeric(valid_review["Volume"], errors="coerce").fillna(0)
        * pd.to_numeric(valid_review["Harga Satuan (Rp)"], errors="coerce").fillna(0)
    )
    st.subheader("Ringkasan Review")
    st.metric(
        "Jumlah jenis kegiatan",
        int(valid_review["Jenis Kegiatan"].astype(str).str.strip().ne("").sum()),
    )
    st.metric(
        "Estimasi total sementara",
        "Rp " + f"{valid_review['Estimasi Total (Rp)'].sum():,.0f}".replace(",", "."),
    )
    review_csv = valid_review.to_csv(index=False).encode("utf-8-sig")
    st.download_button(
        "Unduh hasil review FK/JK (CSV)",
        data=review_csv,
        file_name=f"review_FK_JK_{event_id.strip() or 'event'}.csv",
        mime="text/csv",
    )
else:
    st.caption("Belum ada rincian kegiatan. Tambahkan baris pada tabel untuk memulai review.")


# =====================================================
# VALIDASI V1
# =====================================================

if submitted:
    errors = []

    if not event_id.strip():
        errors.append("Event ID wajib diisi.")

    if not kategori:
        errors.append("Kategori Event wajib dipilih.")

    if not selected_location:
        errors.append("Indeks Lokasi wajib dipilih.")

    if not selected_cc:
        errors.append("Cost Center wajib dipilih.")

    if not selected_account:
        errors.append("Akun Beban Anggaran wajib dipilih.")

    if not lokasi_event.strip():
        errors.append("Lokasi Event wajib diisi.")

    if tanggal_mulai is None:
        errors.append("Tanggal Pelaksanaan wajib diisi.")

    if tanggal_selesai is None:
        errors.append("Tanggal Selesai wajib diisi.")

    if (
        tanggal_mulai is not None
        and tanggal_selesai is not None
        and tanggal_selesai < tanggal_mulai
    ):
        errors.append(
            "Tanggal Selesai tidak boleh lebih awal "
            "daripada Tanggal Pelaksanaan."
        )

    if not uploaded_file:
        errors.append("Dokumen FK/JK wajib diunggah.")

    if not akomodasi:
        errors.append("Akomodasi wajib dipilih.")

    if valid_review.empty:
        errors.append("Minimal satu rincian kegiatan FK/JK harus diisi pada tabel review.")
    else:
        if valid_review["Jenis Kegiatan"].fillna("").astype(str).str.strip().eq("").any():
            errors.append("Jenis Kegiatan pada tabel review tidak boleh kosong.")
        if (pd.to_numeric(valid_review["Volume"], errors="coerce").fillna(0) <= 0).any():
            errors.append("Volume setiap rincian kegiatan harus lebih besar dari 0.")

    if errors:
        st.error("Data belum dapat diproses:")
        for error in errors:
            st.write(f"- {error}")
    else:
        st.success("Input awal lolos validasi.")

        st.info(
            "Data event dan tabel review FK/JK lolos validasi awal. "
            "Perhitungan RAB otomatis dan penerapan norma biaya belum diaktifkan pada versi ini."
        )
