from pathlib import Path
from datetime import date
import io, re, zipfile

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pypdf import PdfReader

st.set_page_config(page_title="AI RAB Generator", page_icon="📊", layout="wide")
MASTER_FILE = Path("Master/Master_RAB_Generator.xlsx")
st.title("AI RAB Generator")
st.caption("Alur: Upload FK/JK → Ekstraksi → Review & validasi norma → Generate Excel")

# ---------- Master data ----------
@st.cache_data
def load_master(file_path, modified_time):
    wb = load_workbook(file_path, data_only=True)
    required = ["TEMPLATE RAB", "INDEKS RAB", "COST CENTER & BEBAN ANGGARAN", "NORMA BIAYA"]
    missing = [s for s in required if s not in wb.sheetnames]
    if missing:
        raise ValueError("Sheet master tidak ditemukan: " + ", ".join(missing))
    norm_ws = wb["NORMA BIAYA"]
    norm_rows = list(norm_ws.iter_rows(values_only=True))
    norms = pd.DataFrame(norm_rows[1:], columns=[str(x).strip() if x is not None else f"Kolom {i+1}" for i, x in enumerate(norm_rows[0])]) if norm_rows else pd.DataFrame()
    norms = norms.dropna(how="all") if not norms.empty else norms
    categories = sorted({str(r[0]).strip() for r in norm_ws.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True) if r[0] is not None and str(r[0]).strip() in {"R1", "R2", "L1", "L2"}})
    ws_i = wb["INDEKS RAB"]
    locations = [{"name": str(r[0]).strip(), "index": r[1]} for r in ws_i.iter_rows(min_row=2, min_col=7, max_col=8, values_only=True) if r[0] is not None and str(r[0]).strip()]
    ws_cc = wb["COST CENTER & BEBAN ANGGARAN"]
    cc = [{"code": str(r[0]).strip(), "description": str(r[1] or "").strip()} for r in ws_cc.iter_rows(min_row=2, min_col=2, max_col=3, values_only=True) if r[0] is not None and str(r[0]).strip()]
    accounts = [{"code": str(r[0]).strip(), "description": str(r[1] or "").strip()} for r in ws_cc.iter_rows(min_row=2, min_col=5, max_col=6, values_only=True) if r[0] is not None and str(r[0]).strip()]
    return categories, locations, cc, accounts, norms

if not MASTER_FILE.exists():
    st.error("File master belum tersedia. Letakkan Master_RAB_Generator.xlsx di folder Master/.")
    st.stop()
try:
    categories, locations, cost_centers, accounts, norms = load_master(str(MASTER_FILE), MASTER_FILE.stat().st_mtime)
except Exception as e:
    st.error(f"Gagal membaca master: {e}")
    st.stop()

# ---------- Extraction helpers ----------
def extract_pdf(file_bytes):
    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)

def extract_image(file_bytes):
    try:
        from PIL import Image
        import pytesseract
        img = Image.open(io.BytesIO(file_bytes))
        return pytesseract.image_to_string(img, lang="ind+eng")
    except Exception:
        return ""

def parse_amount(s):
    s = s.strip().replace("Rp", "").replace("rp", "").replace(" ", "")
    if not s: return None
    # Indonesian format 1.250,50 or 1250,50; plain 1250.00 also handled conservatively
    if "," in s and "." in s: s = s.replace(".", "").replace(",", ".")
    elif "," in s: s = s.replace(".", "").replace(",", ".")
    elif s.count(".") >= 1 and re.search(r"\.\d{3}(\.|$)", s): s = s.replace(".", "")
    try: return float(s)
    except ValueError: return None

def extract_rows(text):
    rows = []
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip(" |\t")
        if len(line) < 4: continue
        # First pass: tabular-looking lines with at least one numeric token
        tokens = re.split(r"\s{2,}|\s\|\s|\t", raw.strip())
        nums = [(i, parse_amount(t)) for i, t in enumerate(tokens) if parse_amount(t) is not None]
        if len(tokens) >= 2 and nums:
            num_i, num = nums[0]
            description = " ".join(tokens[:num_i]).strip(" -:|.")
            if len(description) >= 3 and not re.search(r"^(no|nomor|volume|jumlah|harga|total|uraian|jenis kegiatan)$", description, re.I):
                # Prefer numeric field that resembles volume; default first numeric token
                unit = tokens[num_i + 1] if num_i + 1 < len(tokens) and not parse_amount(tokens[num_i + 1]) else ""
                rows.append({"Jenis Kegiatan": description, "Volume": num, "Satuan": unit, "Harga Satuan": 0.0, "Sumber teks": line})
                continue
        # fallback: capture lines ending in a number, preserving description
        m = re.match(r"^(?:\d+[.)]?\s+)?(.{3,}?)\s+(\d+(?:[.,]\d+)?)\s*(orang|org|hari|paket|unit|kali|pcs|set|jam|malam)?$", line, re.I)
        if m and not re.search(r"total|jumlah|tanggal|halaman", m.group(1), re.I):
            rows.append({"Jenis Kegiatan": m.group(1).strip(" -:|."), "Volume": parse_amount(m.group(2)) or 0, "Satuan": m.group(3) or "", "Harga Satuan": 0.0, "Sumber teks": line})
    # remove duplicates and obvious headings
    seen, cleaned = set(), []
    for r in rows:
        key = (r["Jenis Kegiatan"].lower(), r["Volume"], r["Satuan"].lower())
        if key not in seen and r["Jenis Kegiatan"]:
            seen.add(key); cleaned.append(r)
    return cleaned

def identify_columns(df):
    cols = list(df.columns)
    def find(patterns):
        for c in cols:
            if any(re.search(p, str(c), re.I) for p in patterns): return c
        return None
    return find([r"jenis.*kegiatan", r"uraian", r"nama.*biaya", r"deskripsi", r"komponen", r"jenis biaya", r"kegiatan"]), find([r"norma", r"maksimum", r"harga.*satuan", r"tarif", r"batas.*biaya", r"nilai.*biaya", r"nominal", r"biaya"])

def norm_cap(description, norms_df):
    if norms_df.empty: return None, "Master norma kosong"
    item_col, value_col = identify_columns(norms_df)
    if not item_col or not value_col or item_col == value_col:
        return None, "Kolom item/harga norma belum dapat dikenali otomatis; perlu pemetaan master"
    desc = re.sub(r"[^a-z0-9 ]", " ", str(description).lower())
    desc_words = {w for w in desc.split() if len(w) > 2}
    best = None; best_score = 0
    for _, row in norms_df.iterrows():
        item = str(row.get(item_col, "") or "")
        words = {w for w in re.sub(r"[^a-z0-9 ]", " ", item.lower()).split() if len(w) > 2}
        score = len(desc_words & words) / max(1, len(words))
        val = parse_amount(str(row.get(value_col, ""))) if pd.notna(row.get(value_col)) else None
        if score > best_score and val is not None:
            best, best_score = val, score
    if best is not None and best_score >= 0.5: return best, f"Cocok berdasarkan kata kunci ({best_score:.0%})"
    return None, "Norma tidak cocok otomatis; perlu review manual"

# ---------- Event input ----------
with st.expander("Data event (opsional untuk melengkapi file RAB)", expanded=True):
    c1, c2 = st.columns(2)
    with c1:
        event_id = st.text_input("Event ID", key="event_id")
        category = st.selectbox("Kategori", options=[""] + categories, key="category")
        location_options = [f"{x['name']} — Indeks: {x['index']}" for x in locations]
        selected_location = st.selectbox("Indeks lokasi", options=[""] + location_options, key="location")
        start_date = st.date_input("Tanggal mulai", value=None, key="start_date")
    with c2:
        event_location = st.text_input("Lokasi event", key="event_location")
        cc_options = [f"{x['code']} — {x['description']}" for x in cost_centers]
        selected_cc = st.selectbox("Cost center", options=[""] + cc_options, key="cc")
        account_options = [f"{x['code']} — {x['description']}" for x in accounts]
        selected_account = st.selectbox("Akun beban anggaran", options=[""] + account_options, key="account")
        end_date = st.date_input("Tanggal selesai", value=None, key="end_date")

st.header("1. Upload FK/JK")
uploaded = st.file_uploader("Unggah FK/JK (PDF, PNG, JPG)", type=["pdf", "png", "jpg", "jpeg"])
if uploaded:
    file_bytes = uploaded.getvalue()
    if uploaded.name.lower().endswith(".pdf"):
        try: extracted_text = extract_pdf(file_bytes)
        except Exception as e: extracted_text = ""; st.warning(f"PDF tidak dapat dibaca: {e}")
    else:
        extracted_text = extract_image(file_bytes)
        if not extracted_text:
            st.warning("OCR gambar belum tersedia pada lingkungan ini. Untuk Streamlit Cloud, pasang Tesseract OCR dan paket pytesseract; atau unggah PDF yang memiliki teks selectable.")
    with st.expander("Teks hasil ekstraksi", expanded=False):
        st.text_area("Teks FK/JK", value=extracted_text, height=180, key="extracted_text_view")
    if not extracted_text.strip():
        st.info("Belum ada teks yang bisa diekstrak. PDF hasil scan memerlukan OCR.")
    else:
        st.header("2. Review hasil ekstraksi")
        extracted_rows = extract_rows(extracted_text)
        if not extracted_rows:
            st.warning("Belum ditemukan baris kegiatan dan volume dengan pola yang dikenali. Tambahkan baris secara manual di tabel review.")
            extracted_rows = [{"Jenis Kegiatan":"", "Volume":1.0, "Satuan":"", "Harga Satuan":0.0, "Sumber teks":"Input manual"}]
        base_df = pd.DataFrame(extracted_rows)
        edited = st.data_editor(base_df, num_rows="dynamic", use_container_width=True, key="review_table", column_config={
            "Jenis Kegiatan": st.column_config.TextColumn(required=True),
            "Volume": st.column_config.NumberColumn(min_value=0, step=1, required=True),
            "Satuan": st.column_config.TextColumn(),
            "Harga Satuan": st.column_config.NumberColumn(min_value=0, step=1000, help="Isi atau koreksi harga satuan; norma yang dikenali akan membatasi harga ini."),
            "Sumber teks": st.column_config.TextColumn(disabled=True),
        })
        if st.button("Validasi norma biaya", type="primary"):
            validations = []
            for _, r in edited.iterrows():
                cap, reason = norm_cap(r.get("Jenis Kegiatan", ""), norms)
                price = parse_amount(str(r.get("Harga Satuan", 0))) or 0
                final_price = min(price, cap) if cap is not None and price > 0 else (cap if cap is not None and price <= 0 else price)
                status = "OK" if cap is not None and final_price <= cap else ("Perlu review" if cap is None else "Disesuaikan ke norma")
                validations.append({**r.to_dict(), "Batas Norma": cap if cap is not None else "Tidak terpetakan", "Harga Disetujui": final_price, "Total": (parse_amount(str(r.get("Volume", 0))) or 0) * final_price, "Status Norma": status, "Catatan Validasi": reason})
            st.session_state["validated_rows"] = pd.DataFrame(validations)
            st.session_state["validated_meta"] = {"event_id":event_id, "category":category, "selected_location":selected_location, "event_location":event_location, "selected_cc":selected_cc, "selected_account":selected_account, "start_date":start_date, "end_date":end_date, "filename":uploaded.name}
        if "validated_rows" in st.session_state:
            validated = st.session_state["validated_rows"]
            st.header("3. Hasil validasi norma biaya")
            st.dataframe(validated, use_container_width=True, hide_index=True)
            total = pd.to_numeric(validated["Total"], errors="coerce").fillna(0).sum()
            st.metric("Total estimasi RAB", f"Rp {total:,.0f}".replace(",", "."))
            st.caption("Harga yang belum dapat dipetakan ke norma ditandai 'Perlu review'. Jangan anggap lolos norma sebelum mapping master diverifikasi.")
            if st.button("Generate Excel RAB", type="primary", use_container_width=True):
                output = io.BytesIO()
                wb = load_workbook(MASTER_FILE)
                sheet_name = "HASIL REVIEW FK-JK"
                if sheet_name in wb.sheetnames: del wb[sheet_name]
                ws = wb.create_sheet(sheet_name)
                meta = st.session_state.get("validated_meta", {})
                ws.append(["AI RAB GENERATOR — HASIL REVIEW FK/JK"])
                ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
                ws["A1"].fill = PatternFill("solid", fgColor="17365D")
                ws.append(["Event ID", meta.get("event_id", "")])
                ws.append(["Kategori", meta.get("category", "")])
                ws.append(["Lokasi event", meta.get("event_location", "")])
                ws.append(["Indeks lokasi", meta.get("selected_location", "")])
                ws.append(["Cost center", meta.get("selected_cc", "")])
                ws.append(["Akun beban", meta.get("selected_account", "")])
                ws.append(["Tanggal mulai", str(meta.get("start_date", ""))])
                ws.append(["Tanggal selesai", str(meta.get("end_date", ""))])
                ws.append([])
                headers = list(validated.columns)
                ws.append(headers)
                header_row = ws.max_row
                for cell in ws[header_row]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="4472C4")
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                for row in validated.itertuples(index=False, name=None):
                    ws.append([None if pd.isna(v) else v for v in row])
                for col in ws.columns:
                    letter = col[0].column_letter
                    max_len = max((len(str(c.value or "")) for c in col), default=12)
                    ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)
                ws.freeze_panes = f"A{header_row+1}"
                ws.auto_filter.ref = f"A{header_row}:{ws.cell(ws.max_row, len(headers)).coordinate}"
                ws.append([]); ws.append(["Total estimasi RAB", total])
                wb.save(output); output.seek(0)
                filename = re.sub(r"[^A-Za-z0-9_-]", "_", meta.get("event_id") or "RAB") + "_Review_RAB.xlsx"
                st.download_button("Unduh Excel RAB", data=output.getvalue(), file_name=filename, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
